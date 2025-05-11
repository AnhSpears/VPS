# backend/app.py

import time, json
from collections import deque
from pathlib import Path

from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import JSONResponse
from pydantic import BaseModel
import aiofiles

from config import SHORT_TERM_LIMIT, LOG_FILE, CHAT_HISTORY_FILE
from search_utils import chunk_text, embed_and_index, search

# --- Setup ---
app = FastAPI(title="AI + Document Search Service")

# In-memory buffer
chat_buffer = deque(maxlen=SHORT_TERM_LIMIT)

# Ensure data directories
DATA_DIR   = Path(__file__).parent / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)

# --- Models ---
class ChatRequest(BaseModel):
    text: str
    user_id: str = "default"
class ChatResponse(BaseModel):
    generated: str

# --- Middleware for logging latency & errors ---
@app.middleware("http")
async def log_metrics(request: Request, call_next):
    start = time.perf_counter()
    try:
        response = await call_next(request)
        return response
    except Exception as e:
        # Log error entry
        entry = {
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "path": request.url.path,
            "error": str(e)
        }
        logs = json.loads(LOG_FILE.read_text(encoding="utf-8"))
        logs.append(entry)
        LOG_FILE.write_text(json.dumps(logs, ensure_ascii=False, indent=2), encoding="utf-8")
        raise
    finally:
        elapsed = (time.perf_counter() - start) * 1000
        entry = {
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "path": request.url.path,
            "latency_ms": round(elapsed, 2)
        }
        logs = json.loads(LOG_FILE.read_text(encoding="utf-8"))
        logs.append(entry)
        LOG_FILE.write_text(json.dumps(logs, ensure_ascii=False, indent=2), encoding="utf-8")

# --- Chat endpoint ---
@app.post("/chat", response_model=ChatResponse)
async def chat(req: ChatRequest):
    prompt = req.text.strip()
    if not prompt:
        raise HTTPException(400, "Empty prompt")

    # TODO: Gọi model thực tế
    response_text = f"🤖 AI says: {prompt}"

    # Store to memory
    chat_buffer.append({"prompt": prompt, "response": response_text})
    async with aiofiles.open(CHAT_HISTORY_FILE, "a", encoding="utf-8") as f:
        await f.write(json.dumps({"prompt": prompt, "response": response_text}, ensure_ascii=False) + "\n")

    return {"generated": response_text}

# --- Upload endpoint ---
@app.post("/upload/")
async def upload(file: UploadFile = File(...), field: str = Form(...)):
    """Upload file + field, tự chunk & index."""
    dest = UPLOAD_DIR / file.filename
    async with aiofiles.open(dest, "wb") as out:
        await out.write(await file.read())

    # Extract text
    text = ""
    suffix = file.filename.lower().rsplit(".", 1)[-1]
    if suffix == "pdf":
        import PyPDF2
        reader = PyPDF2.PdfReader(str(dest))
        for p in reader.pages:
            text += p.extract_text() or ""
    elif suffix == "docx":
        import docx
        doc = docx.Document(str(dest))
        for p in doc.paragraphs:
            text += p.text + " "
    elif suffix == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(dest), read_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        text += str(cell) + " "
    else:  # txt
        async with aiofiles.open(dest, "r", encoding="utf-8") as f:
            text = await f.read()

    # Chunk & index
    chunks = chunk_text(text)
    embed_and_index(chunks, file.filename, field)

    return {"status": "indexed", "chunks": len(chunks)}

# --- Search endpoint ---
@app.get("/search/")
async def semantic_search(q: str, k: int = 5):
    """Semantic search trên chunks."""
    results = search(q, top_k=k)
    return {"query": q, "results": results}

# --- Health check ---
@app.get("/health")
async def health():
    return {"status": "ok"}
